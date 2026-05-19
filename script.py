#!/usr/bin/env python3
"""
Fill a class grade Excel file with weekly Code Runner CSV results.

Usage:
    python fill_grades.py <grades.csv> <class.xlsx> <week>

Examples:
    python fill_grades.py week1_1.csv IF-49-01.xlsx 1.1
    python fill_grades.py fungsi.csv IF-49-01.xlsx 2
    python fill_grades.py "array.csv" IF-49-01.xlsx 4

The <week> argument matches the PEKAN label in row 4 of the Excel file.
Accepted forms: "1.1", "1.2", "1", "2", "3", "4", "5", "10", "11-12".
"""

import csv
import sys
from pathlib import Path

from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

HEADER_ROW = 4        # Row in the xlsx containing "KET" / "PEKAN x: ..." labels
DATA_START_ROW = 5    # First student row
NIM_COL = 2           # Column B: NIM
KET_NOT_DONE = "TIDAK MENGERJAKAN"   # Student never started the quiz
KET_NOT_FINISHED = "TIDAK SELESAI"   # Student started but did not submit


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def normalize_week(label: str) -> str:
    """Reduce a 'PEKAN X: ...' label or a user-supplied week to a key like '1.1', '11-12'."""
    if label is None:
        return ""
    s = str(label).strip().upper()
    s = s.replace("PEKAN", "").strip()
    # Drop everything after ':' (the topic name), keep only the week identifier
    s = s.split(":", 1)[0]
    s = s.strip().rstrip(".").strip()
    return s


def find_week_column(ws, week_arg: str) -> int:
    """Locate the score column in HEADER_ROW whose normalized label matches week_arg."""
    target = normalize_week(week_arg)
    if not target:
        raise ValueError(f"Empty week argument: {week_arg!r}")

    matches = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=HEADER_ROW, column=col).value
        if val is None or str(val).upper() == "KET":
            continue
        if normalize_week(val) == target:
            matches.append((col, val))

    if not matches:
        available = [
            str(ws.cell(row=HEADER_ROW, column=c).value)
            for c in range(1, ws.max_column + 1)
            if ws.cell(row=HEADER_ROW, column=c).value
            and str(ws.cell(row=HEADER_ROW, column=c).value).upper() != "KET"
        ]
        raise ValueError(
            f"Could not find week {week_arg!r} in the Excel header.\n"
            f"Available weeks: {available}"
        )
    if len(matches) > 1:
        cols = [m[0] for m in matches]
        raise ValueError(f"Ambiguous week {week_arg!r} matches columns {cols}.")

    return matches[0][0]


def parse_csv(csv_path: Path) -> dict:
    """Return {nim: (state, grade_or_none)} parsed from a Moodle-style grades CSV.

    grade_or_none is a float when the row has a numeric grade, otherwise None.
    Rows without a NIM (e.g. the 'Overall average' footer) are skipped.
    """
    out = {}
    with csv_path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)

        required = {"ID number", "State", "Grade/100.00"}
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(
                f"CSV is missing required columns: {sorted(missing)}.\n"
                f"Found: {reader.fieldnames}"
            )

        for row in reader:
            nim = (row.get("ID number") or "").strip()
            if not nim:
                continue  # skips 'Overall average' and other footer rows
            state = (row.get("State") or "").strip()
            grade_raw = (row.get("Grade/100.00") or "").strip()
            grade = None
            if grade_raw and grade_raw != "-":
                try:
                    grade = float(grade_raw)
                except ValueError:
                    grade = None
            out[nim] = (state, grade)
    return out


def fill_workbook(xlsx_path: Path, csv_data: dict, week_arg: str) -> tuple[int, list[str]]:
    """Fill the workbook in place. Returns (updated_count, list_of_unmatched_nims)."""
    wb = load_workbook(xlsx_path)
    ws = wb.active  # single-sheet file

    score_col = find_week_column(ws, week_arg)
    ket_col = score_col - 1
    ket_label = ws.cell(row=HEADER_ROW, column=ket_col).value
    if ket_label is None or str(ket_label).upper() != "KET":
        raise ValueError(
            f"Expected 'KET' immediately left of the score column "
            f"(col {score_col}), found {ket_label!r}."
        )

    # Build NIM -> row index
    nim_to_row = {}
    for r in range(DATA_START_ROW, ws.max_row + 1):
        nim = ws.cell(row=r, column=NIM_COL).value
        if nim is None:
            continue
        nim_to_row[str(nim).strip()] = r

    updated = 0
    unmatched = []
    for nim, (state, grade) in csv_data.items():
        row = nim_to_row.get(nim)
        if row is None:
            unmatched.append(nim)
            continue

        ket_cell = ws.cell(row=row, column=ket_col)
        score_cell = ws.cell(row=row, column=score_col)

        if state == "Finished" and grade is not None:
            ket_cell.value = None
            score_cell.value = grade
        elif state == "In progress":
            ket_cell.value = KET_NOT_FINISHED
            score_cell.value = 0
        else:
            # Unknown state — be conservative: treat as not done
            ket_cell.value = KET_NOT_DONE
            score_cell.value = 0
        updated += 1

    wb.save(xlsx_path)
    return updated, unmatched


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv: list[str]) -> int:
    if len(argv) != 4:
        print(__doc__)
        return 2

    csv_path = Path(argv[1])
    xlsx_path = Path(argv[2])
    week_arg = argv[3]

    if not csv_path.is_file():
        print(f"Error: CSV not found: {csv_path}", file=sys.stderr)
        return 1
    if not xlsx_path.is_file():
        print(f"Error: XLSX not found: {xlsx_path}", file=sys.stderr)
        return 1

    try:
        csv_data = parse_csv(csv_path)
        if not csv_data:
            print("Error: CSV contains no student rows.", file=sys.stderr)
            return 1
        updated, unmatched = fill_workbook(xlsx_path, csv_data, week_arg)
    except ValueError as e:
        print(f"Error: {e}", file=sys.stderr)
        return 1

    print(f"Updated {updated} student row(s) for week {week_arg!r} in {xlsx_path.name}")
    if unmatched:
        print(f"\nWarning: {len(unmatched)} NIM(s) from the CSV were not found in the XLSX:")
        for nim in unmatched:
            print(f"  - {nim}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))